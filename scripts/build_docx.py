import os, re, sys
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC, OUT = sys.argv[1], sys.argv[2]
_FIG = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'figuras')
FIGS = {'Figure 3.': os.path.join(_FIG, 'Figure3.png'), 'Figure 4.': os.path.join(_FIG, 'Figure4.png')}

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(3)
sec.top_margin = sec.bottom_margin = Cm(2.5)
# continuous line numbering
ln = OxmlElement('w:lnNumType'); ln.set(qn('w:countBy'), '1'); ln.set(qn('w:restart'), 'continuous'); ln.set(qn('w:distance'), '284')
sec._sectPr.append(ln)
# page number in footer (centred)
fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = fp.add_run()
for t, txt in [('begin', None), (None, 'PAGE'), ('end', None)]:
    if t:
        e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), t); r._r.append(e)
    else:
        e = OxmlElement('w:instrText'); e.set(qn('xml:space'), 'preserve'); e.text = txt; r._r.append(e)

st = doc.styles['Normal']; st.font.name = 'Times New Roman'; st.font.size = Pt(12)
st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
pf = st.paragraph_format; pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE; pf.space_after = Pt(6)

def runs(p, text, size=None, bold=None):
    for part in re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*)', text):
        if not part: continue
        b = part.startswith('**'); i = (not b) and part.startswith('*')
        t = part.strip('*') if (b or i) else part
        r = p.add_run(t); r.bold = True if (b or bold) else None; r.italic = True if i else None
        if size: r.font.size = Pt(size)
    return p

def para(text, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=None, bold=None):
    p = doc.add_paragraph(); p.alignment = align; return runs(p, text, size, bold)

def no_lines(p):
    pPr = p._p.get_or_add_pPr(); s = OxmlElement('w:suppressLineNumbers'); pPr.append(s)

lines = open(SRC, encoding='utf-8').read().split('\n')
i = 0
while i < len(lines):
    L = lines[i].rstrip()
    if not L: i += 1; continue
    if L.startswith('@TITLE '):
        para(L[7:], WD_ALIGN_PARAGRAPH.CENTER, 12, True)
    elif L.startswith('@TITLEPT '):
        para(L[9:], WD_ALIGN_PARAGRAPH.CENTER, 12, True)
    elif L.startswith('@AUTHORS '):
        para(L[9:], WD_ALIGN_PARAGRAPH.CENTER)
    elif L.startswith('@AFFIL '):
        p = para(L[7:], WD_ALIGN_PARAGRAPH.LEFT, 10); p.paragraph_format.space_after = Pt(0)
    elif L.startswith('# '):
        p = para(L[2:], WD_ALIGN_PARAGRAPH.LEFT, 12, True); p.paragraph_format.space_before = Pt(12)
    elif L.startswith('## '):
        p = para(L[3:], WD_ALIGN_PARAGRAPH.LEFT, 12, True)
    elif L.startswith('@EQ '):
        para(L[4:], WD_ALIGN_PARAGRAPH.CENTER)
    elif L.startswith('@FIGURE '):
        cap = L[8:]
        key = next((k for k in FIGS if cap.startswith(k)), None)
        if key:
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(FIGS[key], width=Cm(15))
        else:
            para('[' + cap.split('.')[0] + ' – original image from v3.3; labels to be translated into English]', WD_ALIGN_PARAGRAPH.CENTER, 10)
        para(cap, WD_ALIGN_PARAGRAPH.JUSTIFY, 10).runs[0].bold = True
    elif L.startswith('@TABLE '):
        para(L[7:], WD_ALIGN_PARAGRAPH.LEFT, 10).runs[0].bold = True
        rows = []
        i += 1
        while i < len(lines) and lines[i].startswith('|'):
            rows.append([c.strip() for c in lines[i].strip().strip('|').split('|')]); i += 1
        t = doc.add_table(rows=len(rows), cols=len(rows[0]))
        for ri, row in enumerate(rows):
            for ci, c in enumerate(row):
                cell = t.cell(ri, ci); cell.paragraphs[0].paragraph_format.line_spacing = 1.0
                runs(cell.paragraphs[0], c, 9 if len(rows[0]) > 5 else 10, ri == 0)
        # horizontal lines only: top, under header, bottom
        tbl = t._tbl; tblPr = tbl.tblPr; b = OxmlElement('w:tblBorders')
        for edge, v in [('top', 'single'), ('bottom', 'single'), ('insideH', 'nil'), ('insideV', 'nil'), ('left', 'nil'), ('right', 'nil')]:
            e = OxmlElement('w:' + edge); e.set(qn('w:val'), v); e.set(qn('w:sz'), '6'); b.append(e)
        tblPr.append(b)
        for c in t.rows[0].cells:
            tcPr = c._tc.get_or_add_tcPr(); bb = OxmlElement('w:tcBorders'); e = OxmlElement('w:bottom'); e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '6'); bb.append(e); tcPr.append(bb)
        continue
    elif L.startswith('@NOTE '):
        para(L[6:], WD_ALIGN_PARAGRAPH.LEFT, 10)
    else:
        para(L)
    i += 1
doc.save(OUT)
print('saved', OUT)
